import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def add_formatted_text(paragraph, text, bold_all=False):
    """Adds formatted text to a paragraph, handling bold (**) and italics (*)."""
    bold_parts = text.split('**')
    for j, b_part in enumerate(bold_parts):
        italic_parts = b_part.split('*')
        for k, i_part in enumerate(italic_parts):
            if not i_part: continue
            run = paragraph.add_run(i_part)
            if bold_all or (j % 2 == 1):
                run.bold = True
            if k % 2 == 1:
                run.italic = True

def generate_assignment_docx(md_path, docx_path):
    doc = docx.Document()
    
    # Set default style (12pt Times New Roman, Double Spaced)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    paragraph_format = style.paragraph_format
    paragraph_format.line_spacing = 2.0
    
    # Configure Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Metadata
    core_props = doc.core_properties
    core_props.author = "Joel Palk-Ricard"
    core_props.title = "GENS2101 Assignment Final Refined"

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    is_title_page = True
    in_references = False

    for i, line in enumerate(lines):
        line = line.strip()
        
        if not line:
            if not is_title_page:
                doc.add_paragraph()
            continue

        # Handle Title Page
        if is_title_page:
            # Clean formatting markers from title page lines for logic check
            clean_line = line.replace('**', '').replace('*', '')
            
            if line.startswith('## '):
                doc.add_page_break()
                is_title_page = False
                # Continue to body logic below
            elif line.startswith('# '):
                for _ in range(3): doc.add_paragraph()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_formatted_text(p, line[2:], bold_all=True)
                continue
            elif line.startswith('**') or "Professor" in clean_line or "Name:" in clean_line:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_formatted_text(p, line.replace('**', ''))
                
                if "Professor" in clean_line:
                    doc.add_page_break()
                    is_title_page = False
                continue
            elif len(line) > 60:
                doc.add_page_break()
                is_title_page = False
            else:
                continue

        # Handle Body Content
        if not is_title_page:
            if line.startswith('## '):
                if "References" in line:
                    doc.add_page_break()
                    in_references = True
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run("References")
                    run.bold = True
                    continue
                
                p = doc.add_paragraph()
                add_formatted_text(p, line[3:], bold_all=True)
                continue

            if in_references and line.startswith('* '):
                line = line[2:]

            p = doc.add_paragraph()
            if in_references:
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.5)
            
            add_formatted_text(p, line)
        
    doc.save(docx_path)
    print(f"File saved to: {docx_path}")

if __name__ == "__main__":
    ACADEMIC_DIR = "/Users/infinite27/AILCC_PRIME/AI-MasterMind-Alliance/02_Resources/Academics/GENS-2101"
    
    # 1. Vision Assignment
    MD_VISION = os.path.join(ACADEMIC_DIR, "Assignment_1_Vision_Refined.md")
    DOCX_VISION = os.path.join(ACADEMIC_DIR, "Joel_Palk-Ricard_GENS2101_Vision_FINAL.docx")
    print(f"Generating Vision Document...")
    generate_assignment_docx(MD_VISION, DOCX_VISION)
    
    # 2. Project Part 1
    MD_PROJECT = os.path.join(ACADEMIC_DIR, "Personal_Project_Part_1_Refined.md")
    DOCX_PROJECT = os.path.join(ACADEMIC_DIR, "Joel_Palk-Ricard_GENS2101_Project_Part1_FINAL.docx")
    print(f"Generating Project Part 1 Document...")
    generate_assignment_docx(MD_PROJECT, DOCX_PROJECT)
